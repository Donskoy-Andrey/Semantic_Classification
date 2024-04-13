import io

from striprtf.striprtf import rtf_to_text
from PyPDF2 import PdfReader
import openpyxl
from docx import Document


class ParserFile:
    @staticmethod
    async def read_txt(file):
        contents = await file.read()
        return contents.decode('utf-8')

    @staticmethod
    async def read_rtf(file):
        contents = await file.read()
        text_rtf = contents.decode('UTF-8')
        extracted_text = rtf_to_text(text_rtf)
        return extracted_text

    @staticmethod
    async def read_pdf(file):
        contents = await file.read()
        pdf_reader = PdfReader(io.BytesIO(contents))
        extracted_text = ''
        for page in pdf_reader.pages:
            extracted_text += page.extract_text()
        return extracted_text

    @staticmethod
    async def read_xlsx(file):
        contents = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(contents))
        sheet = wb.active
        extracted_text = ''
        for row in sheet.iter_rows(values_only=True):
            extracted_text += " ".join(row)
            extracted_text += " "

        return extracted_text

    @staticmethod
    async def read_docx(file):
        contents = await file.read()
        docx_file = Document(io.BytesIO(contents))
        extracted_text = ""

        for paragraph in docx_file.paragraphs:
            extracted_text += paragraph.text
            extracted_text += " "
        return extracted_text